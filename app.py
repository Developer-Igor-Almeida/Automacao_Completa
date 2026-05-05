import streamlit as st
import subprocess
import shutil
import psutil
import re
import os
import csv
from io import StringIO
from pathlib import Path
from openpyxl import load_workbook
from streamlit_autorefresh import st_autorefresh

st.set_page_config(
    page_title="Automação Tracers + Mind7",
    page_icon="🚗",
    layout="wide"
)

BASE_DIR = Path(__file__).parent

TRACERS = BASE_DIR / "tracers"
MIND7 = BASE_DIR / "mind7"

LOG_DIR = BASE_DIR / "logs"
LOG_DIR.mkdir(exist_ok=True)

LOG_TRACERS = LOG_DIR / "tracers.log"
LOG_MIND7 = LOG_DIR / "mind7.log"

EXCEL_TRACERS = TRACERS / "saida" / "veiculos_tracers.xlsx"
EXCEL_MIND7 = MIND7 / "entrada" / "veiculos_tracers.xlsx"

RESULTADO = MIND7 / "saida" / "veiculos_tracers_com_cpf.xlsx"
CPFS = MIND7 / "saida" / "somente_cpfs.xlsx"
CPFS_CNPJS = MIND7 / "saida" / "cpfs_e_cnpjs.xlsx"
CNPJS = MIND7 / "saida" / "somente_cnpjs.xlsx"

if not TRACERS.exists():
    st.error(f"❌ Pasta do Tracers não encontrada:\n\n{TRACERS}")
    st.stop()

if not MIND7.exists():
    st.error(f"❌ Pasta do Mind7 não encontrada:\n\n{MIND7}")
    st.stop()

(TRACERS / "saida").mkdir(exist_ok=True)
(MIND7 / "entrada").mkdir(exist_ok=True)
(MIND7 / "saida").mkdir(exist_ok=True)

# =========================
# SESSION STATE
# =========================

valores_iniciais = {
    "processo": None,
    "etapa": "",
    "log_atual": None,
    "erro_pausado": False,
    "mensagem_erro_pausado": "",
    "etapa_tracers_ok": False,
    "etapa_envio_ok": False,
    "etapa_chrome_ok": False,
    "etapa_consulta_ok": False,
    "etapa_tracers_erro": False,
    "etapa_envio_erro": False,
    "etapa_chrome_erro": False,
    "etapa_consulta_erro": False,
    "tracers_iniciado": False,
    "consulta_iniciada": False,
}

for chave, valor in valores_iniciais.items():
    if chave not in st.session_state:
        st.session_state[chave] = valor


# =========================
# FUNÇÕES
# =========================

def processo_rodando():
    p = st.session_state.processo
    return p is not None and p.poll() is None


def icone_status(ok, erro):
    if ok:
        return " ✅"
    if erro:
        return " ❌"
    return ""


def ler_log(log_path):
    if not log_path or not Path(log_path).exists():
        return ""
    return Path(log_path).read_text(encoding="utf-8", errors="ignore")


def corrigir_acentos_log(texto):
    if not texto:
        return ""

    correcoes = {
        "execuÃ§Ã£o": "execução",
        "automaÃ§Ã£o": "automação",
        "nÃ£o": "não",
        "estÃ¡": "está",
        "conexÃ£o": "conexão",
        "informaÃ§Ãµes": "informações",
        "veÃ­culo": "veículo",
        "prÃ³xima": "próxima",
        "validaÃ§Ã£o": "validação",
        "possÃ­vel": "possível",
        "cÃ³digo": "código",
        "usuÃ¡rio": "usuário",
        "aplicaÃ§Ã£o": "aplicação",
        "extraÃ§Ã£o": "extração",
    }

    for errado, certo in correcoes.items():
        texto = texto.replace(errado, certo)

    return texto


def detectar_erro_usuario_saiu(log_texto):
    if not log_texto:
        return False

    texto = log_texto.lower()

    erros = [
        "no devices/emulators found",
        "device not found",
        "uiautomator",
        "window_dump",
        "adb.exe",
        "erro ao executar",
        "app_tracers_fechado",
        "não está aberto em primeiro plano",
        "nao esta aberto em primeiro plano",
    ]

    return any(e in texto for e in erros)


def iniciar_processo(cmd, pasta, log_path, etapa):
    with open(log_path, "w", encoding="utf-8") as log:
        log.write(f"Iniciando etapa: {etapa}\n\n")

    log_file = open(log_path, "a", encoding="utf-8")

    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0

    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"

    try:
        p = subprocess.Popen(
            cmd,
            cwd=pasta,
            stdout=log_file,
            stderr=log_file,
            text=True,
            creationflags=flags,
            env=env
        )

        st.session_state.processo = p
        st.session_state.etapa = etapa
        st.session_state.log_atual = log_path

    except Exception as e:
        st.error(f"❌ Erro ao iniciar processo: {e}")


def parar_processo():
    p = st.session_state.processo

    if p and p.poll() is None:
        try:
            processo = psutil.Process(p.pid)
            for filho in processo.children(recursive=True):
                filho.terminate()
            processo.terminate()
        except Exception:
            p.terminate()

    if st.session_state.etapa == "Extração de placas do Tracers":
        st.session_state.etapa_tracers_erro = True

    if st.session_state.etapa == "Consulta de CPFs/CNPJs no Mind7":
        st.session_state.etapa_consulta_erro = True

    st.session_state.processo = None
    st.session_state.etapa = ""


def pausar_por_erro_tracers():
    parar_processo()

    st.session_state.erro_pausado = True
    st.session_state.etapa_tracers_ok = False
    st.session_state.etapa_tracers_erro = True

    st.session_state.mensagem_erro_pausado = (
        "❌ O celular foi desconectado, o aplicativo Tracers foi fechado "
        "ou a tela correta não está aberta.\n\n"
        "Verifique:\n"
        "- Celular conectado\n"
        "- Depuração USB ativa\n"
        "- Aplicativo Tracers aberto\n"
        "- Tela da lista de veículos visível\n\n"
        "Depois clique em CONTINUAR."
    )

    with open(LOG_TRACERS, "a", encoding="utf-8") as log:
        log.write("\n\n==============================\n")
        log.write("ERRO AMIGÁVEL\n")
        log.write("==============================\n")
        log.write("O processo foi pausado porque o celular/app não estava disponível.\n")
        log.write("Verifique o celular e mantenha o aplicativo Tracers aberto na tela correta.\n")


def detectar_finalizacao_processo():
    p = st.session_state.processo

    if p is not None and p.poll() is not None and st.session_state.etapa:
        codigo = p.returncode
        etapa_atual = st.session_state.etapa
        log_texto = ler_log(st.session_state.log_atual)

        if etapa_atual == "Extração de placas do Tracers":
            if detectar_erro_usuario_saiu(log_texto):
                pausar_por_erro_tracers()

            elif codigo == 0 and st.session_state.tracers_iniciado and EXCEL_TRACERS.exists():
                st.session_state.etapa_tracers_ok = True
                st.session_state.etapa_tracers_erro = False

            else:
                st.session_state.etapa_tracers_ok = False
                st.session_state.etapa_tracers_erro = True

        elif etapa_atual == "Consulta de CPFs/CNPJs no Mind7":
            if codigo == 0 and st.session_state.consulta_iniciada and RESULTADO.exists():
                st.session_state.etapa_consulta_ok = True
                st.session_state.etapa_consulta_erro = False
            else:
                st.session_state.etapa_consulta_ok = False
                st.session_state.etapa_consulta_erro = True

        st.session_state.processo = None
        st.session_state.etapa = ""


def monitorar_erro_em_tempo_real():
    if not processo_rodando():
        return

    if st.session_state.etapa != "Extração de placas do Tracers":
        return

    log_texto = ler_log(st.session_state.log_atual)

    if detectar_erro_usuario_saiu(log_texto):
        pausar_por_erro_tracers()


def limpar_logs():
    for log_file in [LOG_TRACERS, LOG_MIND7]:
        try:
            log_file.write_text("", encoding="utf-8")
        except Exception:
            pass

    st.session_state.log_atual = None


def converter_xlsx_para_csv(caminho_xlsx):
    wb = load_workbook(caminho_xlsx)
    ws = wb.active

    output = StringIO()
    writer = csv.writer(output, delimiter=";")

    for row in ws.iter_rows(values_only=True):
        writer.writerow(["" if valor is None else valor for valor in row])

    return output.getvalue().encode("utf-8-sig")


def calcular_progresso(log_texto):
    if not log_texto:
        return 0

    if "Finalizado" in log_texto or "FINALIZADO" in log_texto:
        return 100

    if "Total salvo no Excel" in log_texto:
        return 100

    match_tela = re.findall(r"Tela\s+(\d+)/(\d+)", log_texto)
    if match_tela:
        atual, total = match_tela[-1]
        atual = int(atual)
        total = int(total)
        if total > 0:
            return min(int((atual / total) * 100), 99)

    match_cpf = re.findall(r"Consultando placa:", log_texto)
    if match_cpf:
        return min(len(match_cpf) * 5, 95)

    return 5


def mensagem_amigavel_por_etapa(etapa, log_texto):
    texto = corrigir_acentos_log(log_texto).lower()

    if st.session_state.erro_pausado:
        return st.session_state.mensagem_erro_pausado

    if etapa == "Extração de placas do Tracers":
        if detectar_erro_usuario_saiu(texto):
            return (
                "⚠️ Problema detectado com o celular ou aplicativo Tracers.\n\n"
                "Verifique se o celular está conectado, desbloqueado e com o Tracers aberto."
            )

        if "finalizado" in texto:
            return "✅ Extração de placas finalizada com sucesso."

        return (
            "📱 Extração de placas em andamento.\n\n"
            "Mantenha o celular conectado, desbloqueado e com o aplicativo Tracers aberto."
        )

    if etapa == "Consulta de CPFs/CNPJs no Mind7":
        if "captcha" in texto:
            return (
                "⚠️ O Mind7 solicitou validação CAPTCHA.\n\n"
                "A automação tentará novamente as placas pendentes. "
                "Se necessário, resolva a validação no navegador."
            )

        if "conectado ao chrome" in texto:
            return (
                "🔎 Consulta no Mind7 em andamento.\n\n"
                "Não feche o Chrome e mantenha o usuário logado no Mind7."
            )

        if "finalizado" in texto:
            return "✅ Consulta no Mind7 finalizada com sucesso."

        return (
            "🔎 Preparando consulta no Mind7.\n\n"
            "Verifique se o Chrome está aberto e logado no sistema."
        )

    if "permissionerror" in texto or "permission denied" in texto:
        return (
            "⚠️ Não foi possível salvar o Excel.\n\n"
            "Feche a planilha se ela estiver aberta e tente novamente."
        )

    if "traceback" in texto or "runtimeerror" in texto:
        return (
            "⚠️ A automação encontrou um problema.\n\n"
            "Verifique se os sistemas estão abertos corretamente e tente novamente."
        )

    return ""


# =========================
# MONITORAMENTO
# =========================

detectar_finalizacao_processo()
monitorar_erro_em_tempo_real()

if processo_rodando():
    st_autorefresh(interval=2000, key="refresh_app")


# =========================
# INTERFACE
# =========================

st.markdown("""
<style>
.big-title {
    font-size: 38px;
    font-weight: 800;
    margin-bottom: 5px;
}
.subtitle {
    font-size: 18px;
    color: #666;
    margin-bottom: 20px;
}
</style>
""", unsafe_allow_html=True)

st.markdown(
    '<div class="big-title">🚗 Automação Completa Tracers + Mind7</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="subtitle">Fluxo automatizado: Tracers → Excel → Mind7 → CPF/CNPJ → Resultado Final</div>',
    unsafe_allow_html=True
)

st.divider()

col1, col2, col3, col4 = st.columns(4)

with col1:
    st.metric("Arquivo Tracers", "OK" if EXCEL_TRACERS.exists() else "Pendente")

with col2:
    st.metric("Entrada Mind7", "OK" if EXCEL_MIND7.exists() else "Pendente")

with col3:
    st.metric("Resultado Final", "OK" if RESULTADO.exists() else "Pendente")

with col4:
    st.metric("Status", "Rodando" if processo_rodando() else "Parado")

st.divider()

if processo_rodando():
    st.warning(f"Processo em execução: {st.session_state.etapa}")

    if st.button("🛑 PARAR AUTOMAÇÃO", type="primary"):
        parar_processo()
        st.error("Automação interrompida.")
        st.rerun()
else:
    st.success("Nenhuma automação em execução no momento.")

if st.session_state.erro_pausado:
    st.error(st.session_state.mensagem_erro_pausado)

    if st.button("▶️ Continuar após correção", type="primary"):
        st.session_state.erro_pausado = False
        st.session_state.mensagem_erro_pausado = ""
        st.session_state.etapa_tracers_erro = False
        st.session_state.tracers_iniciado = True

        iniciar_processo(
            ["python", "-u", "extrair_tracers.py"],
            TRACERS,
            LOG_TRACERS,
            "Extração de placas do Tracers"
        )

        st.rerun()

st.divider()

col_esq, col_dir = st.columns(2)

with col_esq:
    st.subheader("1. Extrair placas do Tracers")

    if not (TRACERS / "extrair_tracers.py").exists():
        st.error(f"Arquivo extrair_tracers.py não encontrado em:\n{TRACERS}")
    else:
        texto_botao_tracers = (
            "▶️ Iniciar extração de placas"
            + icone_status(
                st.session_state.etapa_tracers_ok,
                st.session_state.etapa_tracers_erro
            )
        )

        if st.button(texto_botao_tracers, disabled=processo_rodando()):
            st.session_state.etapa_tracers_ok = False
            st.session_state.etapa_tracers_erro = False
            st.session_state.erro_pausado = False
            st.session_state.mensagem_erro_pausado = ""
            st.session_state.tracers_iniciado = True

            iniciar_processo(
                ["python", "-u", "extrair_tracers.py"],
                TRACERS,
                LOG_TRACERS,
                "Extração de placas do Tracers"
            )
            st.rerun()

    st.subheader("2. Enviar Excel para Mind7")

    texto_botao_envio = (
        "📤 Copiar arquivo para Mind7"
        + icone_status(
            st.session_state.etapa_envio_ok,
            st.session_state.etapa_envio_erro
        )
    )

    if EXCEL_TRACERS.exists():
        if st.button(texto_botao_envio, disabled=processo_rodando()):
            try:
                shutil.copy2(EXCEL_TRACERS, EXCEL_MIND7)
                st.session_state.etapa_envio_ok = True
                st.session_state.etapa_envio_erro = False
                st.success("Arquivo copiado para entrada do Mind7.")
                st.rerun()
            except Exception as e:
                st.session_state.etapa_envio_ok = False
                st.session_state.etapa_envio_erro = True
                st.error(f"Erro ao copiar arquivo: {e}")
    else:
        st.info("Execute a extração de placas primeiro.")

with col_dir:
    st.subheader("3. Abrir Chrome Mind7")

    chrome_bat = MIND7 / "abrir_chrome_mind7.bat"

    texto_botao_chrome = (
        "🌐 Abrir Chrome para login"
        + icone_status(
            st.session_state.etapa_chrome_ok,
            st.session_state.etapa_chrome_erro
        )
    )

    if not chrome_bat.exists():
        st.error(f"Arquivo abrir_chrome_mind7.bat não encontrado em:\n{chrome_bat}")
    else:
        if st.button(texto_botao_chrome, disabled=processo_rodando()):
            try:
                subprocess.Popen(
                    ["cmd", "/c", str(chrome_bat)],
                    cwd=str(MIND7),
                    shell=False
                )
                st.session_state.etapa_chrome_ok = True
                st.session_state.etapa_chrome_erro = False
                st.success("Chrome aberto. Faça login no Mind7.")
                st.rerun()
            except Exception as e:
                st.session_state.etapa_chrome_ok = False
                st.session_state.etapa_chrome_erro = True
                st.error(f"Erro ao abrir Chrome: {e}")

    st.subheader("4. Buscar CPFs/CNPJs")

    texto_botao_consulta = (
        "🔎 Iniciar consulta no Mind7"
        + icone_status(
            st.session_state.etapa_consulta_ok,
            st.session_state.etapa_consulta_erro
        )
    )

    if not (MIND7 / "consultar_mind7.py").exists():
        st.error(f"Arquivo consultar_mind7.py não encontrado em:\n{MIND7}")
    elif EXCEL_MIND7.exists():
        if st.button(texto_botao_consulta, disabled=processo_rodando()):
            st.session_state.etapa_consulta_ok = False
            st.session_state.etapa_consulta_erro = False
            st.session_state.consulta_iniciada = True

            iniciar_processo(
                ["python", "-u", "consultar_mind7.py"],
                MIND7,
                LOG_MIND7,
                "Consulta de CPFs/CNPJs no Mind7"
            )
            st.rerun()
    else:
        st.info("Copie o Excel para o Mind7 primeiro.")

st.divider()

st.subheader("📊 Progresso em tempo real")

log_texto = ler_log(st.session_state.log_atual)
progresso = calcular_progresso(log_texto)

st.progress(progresso / 100)
st.write(f"Progresso estimado: **{progresso}%**")

st.subheader("📜 Logs organizados")

col_log1, col_log2, col_log3 = st.columns([1, 1, 4])

with col_log1:
    if st.button("🧹 Limpar logs", disabled=processo_rodando()):
        limpar_logs()
        st.success("Logs limpos com sucesso!")
        st.rerun()

with col_log2:
    if st.button("🔄 Atualizar logs"):
        st.rerun()

with st.expander("Status da execução", expanded=True):
    log_texto = corrigir_acentos_log(ler_log(st.session_state.log_atual))
    mensagem = mensagem_amigavel_por_etapa(st.session_state.etapa, log_texto)

    if mensagem:
        st.info(mensagem)
    elif log_texto:
        st.success("Processo em andamento ou finalizado sem erro crítico.")
    else:
        st.info("Nenhum processo iniciado ainda.")

with st.expander("Logs técnicos para suporte", expanded=False):
    log_texto = corrigir_acentos_log(ler_log(st.session_state.log_atual))

    if log_texto:
        st.code(log_texto[-8000:])
    else:
        st.info("Nenhum log técnico disponível.")

st.divider()

st.subheader("📥 Baixar arquivo")

arquivos_disponiveis = {
    "Resultado completo - Excel": RESULTADO,
    "Somente CPFs - Excel": CPFS,
    "CPFs + CNPJs - Excel": CPFS_CNPJS,
    "Somente CNPJs - Excel": CNPJS,
    "Resultado completo - CSV": RESULTADO,
    "Somente CPFs - CSV": CPFS,
    "CPFs + CNPJs - CSV": CPFS_CNPJS,
    "Somente CNPJs - CSV": CNPJS,
}

opcao_download = st.selectbox(
    "Escolha o arquivo para baixar:",
    list(arquivos_disponiveis.keys())
)

arquivo_escolhido = arquivos_disponiveis[opcao_download]

if not arquivo_escolhido.exists():
    st.info("Esse arquivo ainda não foi gerado.")
else:
    if "CSV" in opcao_download:
        dados = converter_xlsx_para_csv(arquivo_escolhido)
        nome_arquivo = arquivo_escolhido.stem + ".csv"
        mime = "text/csv"
    else:
        dados = arquivo_escolhido.read_bytes()
        nome_arquivo = arquivo_escolhido.name
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    st.download_button(
        "📥 Baixar arquivo selecionado",
        data=dados,
        file_name=nome_arquivo,
        mime=mime,
        type="primary"
    )