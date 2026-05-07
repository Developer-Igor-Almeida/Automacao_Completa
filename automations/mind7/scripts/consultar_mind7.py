"""
Entrada principal da automação Mind7.

Responsável por:
- carregar o Excel de entrada
- conectar ao Chrome
- orquestrar consultas
- salvar os resultados finais
"""

import time
from openpyxl import load_workbook
from automations.mind7.config.settings import (DOCUMENT_COLUMN_NAME,PLATE_COLUMN_NAME, CAPTCHA_SUCCESS_WAIT_SECONDS)
from automations.mind7.constants.messages import (CHROME_CONNECTED_MESSAGE, DOCUMENT_FOUND_MESSAGE,DOCUMENT_NOT_FOUND_MESSAGE, DOCUMENT_NOT_FOUND_RETRY_MESSAGE,
    FINAL_RESULT_SAVED_MESSAGE, MIND7_HEADER_SEPARATOR, MIND7_TITLE,
    PROCESS_FINISHED_MESSAGE, QUERY_ERROR_MESSAGE, SAVING_FINAL_RESULT_MESSAGE, CAPTCHA_DETECTED_WAIT_MESSAGE,
)
from automations.mind7.constants.query import (CAPTCHA_RESULT,DOCUMENT_NOT_FOUND_VALUE,EMPTY_DOCUMENT_VALUE,QUERY_ERROR_VALUE,)
from automations.mind7.services.chrome_service import ( connect_to_mind7_page,)
from automations.mind7.services.document_service import (normalize_plate,)
from automations.mind7.services.excel_service import (ensure_column, find_column, save_filtered_documents,)
from automations.mind7.services.mind7_service import (query_plate,)
from automations.mind7.services.retry_service import (preventive_pause,retry_captcha_plates,)
from backend.core.paths import (FINAL_RESULT_FILE,MIND7_INPUT_EXCEL_FILE,)

def main() -> None:
    _print_header()
    workbook = load_workbook(MIND7_INPUT_EXCEL_FILE)
    worksheet = workbook.active
    plate_column = _get_plate_column(worksheet)
    document_column = ensure_column(worksheet, DOCUMENT_COLUMN_NAME)
    playwright, browser, page = connect_to_mind7_page()
    print(CHROME_CONNECTED_MESSAGE, flush=True)
    try:
        retry_rows = _process_plates(page=page, worksheet=worksheet, plate_column=plate_column, document_column=document_column)
        if retry_rows:
            retry_captcha_plates(page=page, worksheet=worksheet, plate_column=plate_column, document_column=document_column, captcha_rows=retry_rows)
        _save_results(workbook=workbook, worksheet=worksheet, document_column=document_column)
    finally:
        time.sleep(2)
        browser.close()
        playwright.stop()

def _print_header() -> None:
    print(MIND7_HEADER_SEPARATOR, flush=True)
    print(MIND7_TITLE, flush=True)
    print(MIND7_HEADER_SEPARATOR, flush=True)

def _get_plate_column(worksheet) -> int:
    plate_column = find_column(worksheet, PLATE_COLUMN_NAME)
    if not plate_column:
        raise ValueError(f"Coluna '{PLATE_COLUMN_NAME}' não encontrada.")
    return plate_column

def _process_plates(page, worksheet, plate_column: int, document_column: int) -> list[int]:
    captcha_rows = []
    retry_not_found_rows = []
    completed_queries = 0
    for row in range(2, worksheet.max_row + 1):
        plate = _get_plate_from_row(worksheet=worksheet, row=row, plate_column=plate_column)
        if not plate:
            continue
        _print_query_progress(row=row, total_rows=worksheet.max_row, plate=plate)
        try:
            document = query_plate(page, plate)
            result_status = _handle_query_result(worksheet=worksheet, row=row, document_column=document_column, plate=plate, document=document, captcha_rows=captcha_rows, retry_not_found_rows=retry_not_found_rows)
            if result_status == CAPTCHA_RESULT:
                _wait_after_captcha_success()
                captcha_rows.extend(_get_remaining_rows(current_row=row, max_row=worksheet.max_row))
                break
            completed_queries += 1
            preventive_pause(completed_queries)
        except Exception as error:
            print(QUERY_ERROR_MESSAGE.format(plate=plate, error=error), flush=True)
            _set_document_value(worksheet=worksheet, row=row, document_column=document_column, value=QUERY_ERROR_VALUE)
    return captcha_rows + retry_not_found_rows

def _get_plate_from_row(worksheet, row: int, plate_column: int) -> str:
    original_plate = worksheet.cell(row=row, column=plate_column).value
    return normalize_plate(original_plate)

def _print_query_progress(row: int, total_rows: int, plate: str) -> None:
    print(f"\n[{row - 1}/{total_rows - 1}] Consultando placa: {plate}", flush=True)

def _handle_query_result(worksheet, row: int, document_column: int, plate: str, document: str, captcha_rows: list[int], retry_not_found_rows: list[int]) -> str:
    if document == CAPTCHA_RESULT:
        _handle_captcha_result(worksheet=worksheet, row=row, document_column=document_column, captcha_rows=captcha_rows)
        return CAPTCHA_RESULT
    if document:
        print(DOCUMENT_FOUND_MESSAGE.format(document=document), flush=True)
        _set_document_value(worksheet=worksheet, row=row, document_column=document_column, value=document)
        return "OK"
    print(DOCUMENT_NOT_FOUND_RETRY_MESSAGE.format(plate=plate), flush=True)
    _set_document_value(worksheet=worksheet, row=row, document_column=document_column, value=EMPTY_DOCUMENT_VALUE)
    retry_not_found_rows.append(row)
    return "OK"

def _handle_captcha_result(worksheet, row: int, document_column: int, captcha_rows: list[int]) -> None:
    _set_document_value(worksheet=worksheet, row=row, document_column=document_column, value=EMPTY_DOCUMENT_VALUE)
    captcha_rows.append(row)

def _set_document_value(worksheet, row: int, document_column: int, value: str) -> None:
    worksheet.cell(row=row, column=document_column).value = value

def _save_results(workbook, worksheet, document_column: int) -> None:
    print(SAVING_FINAL_RESULT_MESSAGE, flush=True)
    workbook.save(FINAL_RESULT_FILE)
    save_filtered_documents(worksheet, document_column)
    print(FINAL_RESULT_SAVED_MESSAGE.format(file_path=FINAL_RESULT_FILE), flush=True)
    print(PROCESS_FINISHED_MESSAGE, flush=True)
    
def _get_remaining_rows(current_row: int,max_row: int,) -> list[int]:
    return list(range(current_row + 1, max_row + 1,))

def _wait_after_captcha_success() -> None:
    print(CAPTCHA_DETECTED_WAIT_MESSAGE.format(seconds=CAPTCHA_SUCCESS_WAIT_SECONDS), flush=True)
    time.sleep(CAPTCHA_SUCCESS_WAIT_SECONDS)

if __name__ == "__main__":
    main()