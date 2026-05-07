"""
Serviços responsáveis por leitura, escrita e exportação de Excel do Mind7.
"""

from pathlib import Path

from openpyxl import Workbook

from automations.mind7.services.document_service import (
    format_cnpj,
    format_cpf,
    get_document_type,
    only_numbers,
)
from backend.core.paths import (
    CNPJ_ONLY_FILE,
    CPF_CNPJ_FILE,
    CPF_ONLY_FILE,
)


DOCUMENTS_OUTPUT_CONFIG = {
    "cpf": {
        "file": CPF_ONLY_FILE,
        "sheet_title": "CPFs",
        "headers": ["CPF"],
        "types": {"CPF"},
    },
    "cpf_cnpj": {
        "file": CPF_CNPJ_FILE,
        "sheet_title": "CPFs e CNPJs",
        "headers": ["Documento", "Tipo"],
        "types": {"CPF", "CNPJ"},
    },
    "cnpj": {
        "file": CNPJ_ONLY_FILE,
        "sheet_title": "CNPJs",
        "headers": ["CNPJ"],
        "types": {"CNPJ"},
    },
}


IGNORED_DOCUMENT_VALUES = {
    "não encontrado",
    "nao encontrado",
    "erro",
    "captcha",
}


def find_column(worksheet, column_name: str) -> int | None:
    for column in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=column).value

        if value and str(value).strip().lower() == column_name.lower():
            return column

    return None


def ensure_column(worksheet, column_name: str) -> int:
    column = find_column(worksheet, column_name)

    if column:
        return column

    new_column = worksheet.max_column + 1
    worksheet.cell(row=1, column=new_column).value = column_name

    return new_column


def save_filtered_documents(
    worksheet,
    document_column: int,
) -> None:
    seen_documents = {
        "cpf": set(),
        "cpf_cnpj": set(),
        "cnpj": set(),
    }

    workbooks = _create_output_workbooks()

    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=document_column).value

        if not value:
            continue

        value = str(value).strip()

        if value.lower() in IGNORED_DOCUMENT_VALUES:
            continue

        document_type = get_document_type(value)
        document_numbers = only_numbers(value)

        if not document_type:
            continue

        formatted_document = (
            format_cpf(value)
            if document_type == "CPF"
            else format_cnpj(value)
        )

        _append_document_to_outputs(
            document_type=document_type,
            document_numbers=document_numbers,
            formatted_document=formatted_document,
            seen_documents=seen_documents,
            workbooks=workbooks,
        )

    _save_output_workbooks(workbooks)


def _create_output_workbooks() -> dict:
    workbooks = {}

    for key, config in DOCUMENTS_OUTPUT_CONFIG.items():
        workbook = Workbook()
        worksheet = workbook.active

        worksheet.title = config["sheet_title"]
        worksheet.append(config["headers"])
        worksheet.column_dimensions["A"].width = 24

        if len(config["headers"]) > 1:
            worksheet.column_dimensions["B"].width = 12

        workbooks[key] = (workbook, worksheet)

    return workbooks


def _append_document_to_outputs(
    document_type: str,
    document_numbers: str,
    formatted_document: str,
    seen_documents: dict,
    workbooks: dict,
) -> None:
    for key, config in DOCUMENTS_OUTPUT_CONFIG.items():
        if document_type not in config["types"]:
            continue

        if document_numbers in seen_documents[key]:
            continue

        seen_documents[key].add(document_numbers)

        _, worksheet = workbooks[key]

        if key == "cpf_cnpj":
            worksheet.append([formatted_document, document_type])
        else:
            worksheet.append([formatted_document])


def _save_output_workbooks(workbooks: dict) -> None:
    for key, config in DOCUMENTS_OUTPUT_CONFIG.items():
        workbook, _ = workbooks[key]
        workbook.save(config["file"])