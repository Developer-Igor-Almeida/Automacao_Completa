"""
Utilitários para manipulação e conversão de arquivos.

Responsável por:
- converter arquivos Excel para CSV
- normalizar dados exportados
"""

import csv
from io import StringIO
from pathlib import Path
from openpyxl import load_workbook
from backend.constants.encoding import DEFAULT_ENCODING

def convert_excel_to_csv_bytes(excel_file_path: Path) -> bytes:
    workbook = load_workbook(excel_file_path,data_only=True,)
    worksheet = workbook.active
    output = StringIO()
    csv_writer = csv.writer(output,delimiter=";",)

    for row in worksheet.iter_rows(values_only=True):
        csv_writer.writerow(_normalize_row(row))

    return output.getvalue().encode(f"{DEFAULT_ENCODING}-sig")

def _normalize_row(row: tuple) -> list:
    return [
        "" if value is None else value
        for value in row
    ]